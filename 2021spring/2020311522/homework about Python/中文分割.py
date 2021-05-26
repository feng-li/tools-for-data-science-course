import docx
#获取文档
def getText(filename):
    doc = docx.Document(filename)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    return '\n'.join(fullText)

text = getText('云雀叫了一整天节选.docx')

#文档文字切分
import jieba
words = jieba.cut(text)
# print('/'.join(words))

#词汇词性分类
import jieba.posseg as pseg
result = pseg.cut(text)

# for w in result:
#       print(w.word,w.flag)


#文本转成数据 合并各类词性
import xlwt
import codecs

input_txt = '数据存放.txt'
output_excel = '数据整合.xls'
sheetName = 'Sheet1'
start_row = 0
start_col = 0

wb = xlwt.Workbook(encoding='utf-8')
ws = wb.add_sheet(sheetName)

f = open(input_txt, encoding='utf-8')

row_excel = start_row
for line in f:
    line = line.strip('\n')
    line = line.split(' ')

    # print(line)

    col_excel = start_col
    len_line = len(line)
    for j in range(len_line):
        # print(line[j])
        ws.write(row_excel, col_excel, line[j])
        col_excel += 1
        wb.save(output_excel)

    row_excel += 1

f.close

#缓存文档中数据导入pandas 进行建表
import pandas as pd

output_excel = '数据整合.xls'
feature = pd.read_excel(output_excel, usecols=[0,1])
print(feature)

#词性数据统计
m =feature['x'].value_counts()
print(m)
data = pd.DataFrame(m)#出现次数
print(data)

#绘制成表
import numpy as np
import matplotlib.pyplot as plt
data.plot(kind = 'bar')
